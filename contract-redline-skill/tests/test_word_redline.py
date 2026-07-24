from __future__ import annotations

import json
import subprocess
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import RGBColor


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "word_redline.py"


class WordRedlineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="word-redline-tests-")
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def make_docx(
        self,
        name: str,
        text: str,
        *,
        creator: str = "原作者",
        color: RGBColor | None = None,
    ) -> Path:
        path = self.root / name
        document = Document()
        document.core_properties.author = creator
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text)
        if color is not None:
            run.font.color.rgb = color
        document.save(path)
        return path

    def apply(self, plan: dict) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            ["python3", str(SCRIPT), "apply"],
            input=json.dumps(plan, ensure_ascii=False),
            capture_output=True,
            text=True,
            check=False,
        )

    def command(self, *args: str) -> dict:
        result = subprocess.run(
            ["python3", str(SCRIPT), *args],
            capture_output=True,
            text=True,
            check=True,
        )
        return json.loads(result.stdout)

    def read_part(self, path: Path, part: str) -> str:
        with ZipFile(path) as archive:
            return archive.read(part).decode("utf-8")

    def test_default_is_transactional_in_place_without_extra_docx(self) -> None:
        source = self.make_docx("合同.docx", "甲方应付款。")

        result = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "replace_text",
                        "locator": {"paragraph_index": 1},
                        "old_text": "应",
                        "new_text": "应当",
                    }
                ],
            }
        )

        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        payload = json.loads(result.stdout)
        self.assertEqual(payload["mode"], "in_place")
        self.assertEqual(payload["output"], str(source))
        self.assertEqual(
            sorted(path.name for path in self.root.glob("*.docx")),
            ["合同.docx"],
        )
        self.assertTrue(payload["verify"]["has_revisions"])

    def test_failed_plan_does_not_modify_source_or_leave_temp_docx(self) -> None:
        source = self.make_docx("合同.docx", "甲方应付款。")
        original_bytes = source.read_bytes()

        result = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "replace_text",
                        "locator": {"paragraph_index": 1},
                        "old_text": "应",
                        "new_text": "应当",
                    },
                    {
                        "type": "replace_text",
                        "locator": {"paragraph_index": 1},
                        "old_text": "不存在",
                        "new_text": "其他",
                    },
                ],
            }
        )

        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(source.read_bytes(), original_bytes)
        self.assertEqual(
            sorted(path.name for path in self.root.glob("*.docx")),
            ["合同.docx"],
        )

    def test_timestamps_are_contextual_ordered_and_not_future(self) -> None:
        source = self.make_docx("合同.docx", "甲方付款并在完成验收后开具合法有效的发票。")

        result = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "insert_after",
                        "locator": {"paragraph_index": 1},
                        "anchor_text": "甲方",
                        "text": "应当",
                    },
                    {
                        "type": "insert_after",
                        "locator": {"paragraph_index": 1},
                        "anchor_text": "付款",
                        "text": "并提供付款凭证",
                    },
                    {
                        "type": "add_comment",
                        "locator": {"paragraph_index": 1},
                        "target_text": "合法有效的发票",
                        "comment_text": "请结合税务要求确认发票类型及开具期限。",
                    },
                ],
            }
        )

        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        payload = json.loads(result.stdout)
        timestamps = [
            datetime.fromisoformat(value)
            for value in payload["verify"]["timestamp_samples"]
        ]
        self.assertEqual(len(timestamps), 3)
        self.assertEqual(timestamps, sorted(timestamps))
        self.assertLessEqual(timestamps[-1], datetime.now(timestamps[-1].tzinfo))
        gaps = [
            int((timestamps[index] - timestamps[index - 1]).total_seconds())
            for index in range(1, len(timestamps))
        ]
        self.assertGreater(len(set(gaps)), 1)

    def test_leading_spaces_do_not_shift_insert_anchor(self) -> None:
        source = self.make_docx("合同.docx", "  甲方付款。")

        result = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "insert_after",
                        "locator": {"paragraph_index": 1},
                        "anchor_text": "甲方",
                        "text": "应当",
                    }
                ],
            }
        )

        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        inspected = self.command(
            "inspect",
            "--source",
            str(source),
            "--start",
            "1",
            "--limit",
            "1",
        )
        self.assertEqual(inspected["paragraphs"][0]["text"], "  甲方应当付款。")

    def test_existing_revisions_can_be_preserved_while_editing_plain_text(self) -> None:
        source = self.make_docx("合同.docx", "甲方应付款。")
        first = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "replace_text",
                        "locator": {"paragraph_index": 1},
                        "old_text": "应",
                        "new_text": "须",
                    }
                ],
            }
        )
        self.assertEqual(first.returncode, 0, first.stdout + first.stderr)

        second = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "replace_text",
                        "locator": {"paragraph_index": 1},
                        "old_text": "付款",
                        "new_text": "结算",
                    }
                ],
            }
        )
        self.assertEqual(second.returncode, 0, second.stdout + second.stderr)
        payload = json.loads(second.stdout)
        self.assertEqual(payload["verify"]["new_revisions"], 2)
        self.assertEqual(payload["verify"]["total_revisions"], 4)
        self.assertTrue(payload["verify"]["revision_ids_unique"])
        self.assertEqual(
            sorted(path.name for path in self.root.glob("*.docx")),
            ["合同.docx"],
        )

    def test_creator_and_font_color_are_preserved(self) -> None:
        source = self.make_docx(
            "合同.docx",
            "红色条款",
            creator="原作者",
            color=RGBColor(255, 0, 0),
        )

        result = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "replace_text",
                        "locator": {"paragraph_index": 1},
                        "old_text": "条款",
                        "new_text": "约定",
                    }
                ],
            }
        )

        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        document_xml = self.read_part(source, "word/document.xml")
        core_xml = self.read_part(source, "docProps/core.xml")
        self.assertIn('w:val="FF0000"', document_xml)
        self.assertIn("<dc:creator>原作者</dc:creator>", core_xml)
        self.assertIn("<cp:lastModifiedBy>审阅人</cp:lastModifiedBy>", core_xml)

    def test_inspect_and_verify_report_structural_risks(self) -> None:
        source = self.root / "表格合同.docx"
        document = Document()
        document.add_paragraph("服务内容")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "验收标准"
        document.save(source)

        inspected = self.command(
            "inspect",
            "--source",
            str(source),
            "--start",
            "1",
            "--limit",
            "0",
        )
        self.assertEqual(inspected["structure"]["table_count"], 1)
        table_paragraphs = [
            paragraph
            for paragraph in inspected["paragraphs"]
            if paragraph["container"] == "table"
        ]
        self.assertTrue(table_paragraphs)

        result = self.apply(
            {
                "source": str(source),
                "operations": [
                    {
                        "type": "add_comment",
                        "locator": {"paragraph_index": table_paragraphs[0]["index"]},
                        "target_text": "验收标准",
                        "comment_text": "请确认验收标准是否完整。",
                    }
                ],
            }
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        verified = self.command("verify", "--source", str(source))
        self.assertTrue(verified["comment_links_ok"])
        self.assertTrue(verified["zip_integrity_ok"])
        self.assertTrue(verified["xml_integrity_ok"])


if __name__ == "__main__":
    unittest.main()
