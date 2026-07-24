from __future__ import annotations

import hashlib
import json
import subprocess
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.shared import RGBColor


ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "scripts" / "contract_pipeline.py"
REDLINE = ROOT / "scripts" / "word_redline.py"


class ContractPipelineTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory(prefix="contract-pipeline-tests-")
        self.evidence_temp_dir = tempfile.TemporaryDirectory(prefix="contract-evidence-tests-")
        self.root = Path(self.temp_dir.name)
        self.evidence_root = Path(self.evidence_temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()
        self.evidence_temp_dir.cleanup()

    def make_docx(self, name: str, paragraphs: list[str]) -> Path:
        path = self.root / name
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        document.save(path)
        return path

    def make_mixed_format_docx(self, name: str) -> Path:
        path = self.root / name
        document = Document()
        document.add_paragraph("可安全修改的付款条款。")
        paragraph = document.add_paragraph()
        paragraph.add_run("混合").font.color.rgb = RGBColor(255, 0, 0)
        paragraph.add_run("格式条款。")
        document.add_paragraph("另一安全条款。")
        document.save(path)
        return path

    def run_json(self, script: Path, *args: str, payload: dict | None = None):
        result = subprocess.run(
            ["python3", str(script), *args],
            input=json.dumps(payload, ensure_ascii=False) if payload is not None else None,
            capture_output=True,
            text=True,
            check=False,
        )
        return result, json.loads(result.stdout)

    def snapshot(self, source: Path) -> dict:
        result, payload = self.run_json(PIPELINE, "snapshot", "--source", str(source))
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return payload

    def intent_payload(self, source: Path, snapshot: dict, intents: list[dict]) -> dict:
        return {
            "source": str(source),
            "source_sha256": snapshot["source_sha256"],
            "snapshot_id": snapshot["snapshot_id"],
            "intents": intents,
        }

    def edit_intent(
        self,
        snapshot: dict,
        paragraph_index: int,
        proposed_text: str,
        *,
        intent_id: str,
        dependency_group: str | None = None,
    ) -> dict:
        block = snapshot["paragraphs"][paragraph_index - 1]
        return {
            "intent_id": intent_id,
            "anchor": {"block_id": block["block_id"], "source_quote": block["text"][:8]},
            "risk": "测试风险",
            "desired_effect": "测试法律效果",
            "treatment": "edit",
            "proposed_text": proposed_text,
            "dependency_group": dependency_group or intent_id,
        }

    def test_snapshot_produces_stable_parser_owned_block_ids(self) -> None:
        source = self.make_docx("合同.docx", ["买卖合同", "第一条 标的", "设备一台。"])
        first = self.snapshot(source)
        second = self.snapshot(source)
        self.assertEqual(first["schema_version"], "pan001-document-snapshot/v2")
        self.assertEqual(first["snapshot_id"], second["snapshot_id"])
        self.assertEqual(
            [item["block_id"] for item in first["paragraphs"]],
            [item["block_id"] for item in second["paragraphs"]],
        )
        self.assertTrue(all(item["block_id"].startswith("body:p") for item in first["paragraphs"]))

    def test_prepare_renders_multi_diff_paragraph_once_and_preserves_source(self) -> None:
        original = "甲方验证合格后方可付款"
        proposed = "甲方应在收到完整付款资料后及时核验"
        source = self.make_docx("合同.docx", [original])
        original_bytes = source.read_bytes()
        snapshot = self.snapshot(source)
        plan_path = self.evidence_root / "prepared.json"
        evidence = self.evidence_root / "multi-diff"
        payload = self.intent_payload(
            source,
            snapshot,
            [self.edit_intent(snapshot, 1, proposed, intent_id="PAYMENT-01")],
        )
        result, prepared = self.run_json(
            PIPELINE,
            "prepare",
            "--output-plan",
            str(plan_path),
            "--evidence-dir",
            str(evidence),
            payload=payload,
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(prepared["stage"], "prepared")
        self.assertEqual(prepared["operation_count"], 1)
        self.assertEqual(source.read_bytes(), original_bytes)
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        self.assertEqual(plan["operations"][0]["type"], "render_block")
        self.assertEqual(plan["operations"][0]["locator"]["block_id"], snapshot["paragraphs"][0]["block_id"])
        self.assertTrue((evidence / "01-intents.json").exists())
        self.assertTrue((evidence / "04-prepared-plan.json").exists())

        applied, applied_payload = self.run_json(REDLINE, "apply", payload=plan)
        self.assertEqual(applied.returncode, 0, applied.stdout + applied.stderr)
        self.assertTrue(applied_payload["verify"]["has_revisions"])
        _, inspected = self.run_json(REDLINE, "inspect", "--source", str(source), "--limit", "0")
        self.assertEqual(inspected["paragraphs"][0]["text"], proposed)
        self.assertEqual(inspected["paragraphs"][0]["rejected_text"], original)

    def test_source_hash_and_snapshot_block_stale_intents(self) -> None:
        source = self.make_docx("合同.docx", ["甲方付款。"])
        snapshot = self.snapshot(source)
        document = Document(source)
        document.paragraphs[0].text = "甲方暂不付款。"
        document.save(source)
        payload = self.intent_payload(
            source,
            snapshot,
            [self.edit_intent(snapshot, 1, "甲方应付款。", intent_id="PAY-01")],
        )
        result, error = self.run_json(PIPELINE, "prepare", payload=payload)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn(error["error_code"], {"SNAPSHOT_CHANGED", "SOURCE_CHANGED"})

    def test_source_quote_must_be_continuous_text_in_anchored_block(self) -> None:
        source = self.make_docx("合同.docx", ["甲方应在验收后付款。"])
        snapshot = self.snapshot(source)
        intent = self.edit_intent(snapshot, 1, "甲方应及时付款。", intent_id="PAY-01")
        intent["anchor"]["source_quote"] = "合同中不存在的文字"
        result, error = self.run_json(
            PIPELINE, "prepare", payload=self.intent_payload(source, snapshot, [intent])
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(error["error_code"], "SOURCE_QUOTE_MISMATCH")

    def test_comment_uses_same_stable_block_anchor(self) -> None:
        source = self.make_docx("合同.docx", ["付款期限为30日。"])
        snapshot = self.snapshot(source)
        block = snapshot["paragraphs"][0]
        intent = {
            "intent_id": "COMMENT-01",
            "anchor": {"block_id": block["block_id"], "source_quote": "30日"},
            "risk": "付款期限属于商业参数",
            "desired_effect": "提示用户确认",
            "treatment": "comment",
            "comment_text": "请确认付款期限是否符合商业安排。",
        }
        plan_path = self.evidence_root / "comment.json"
        result, prepared = self.run_json(
            PIPELINE,
            "prepare",
            "--output-plan",
            str(plan_path),
            payload=self.intent_payload(source, snapshot, [intent]),
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(prepared["operation_count"], 1)
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        self.assertEqual(plan["operations"][0]["locator"]["block_id"], block["block_id"])
        applied, payload = self.run_json(REDLINE, "apply", payload=plan)
        self.assertEqual(applied.returncode, 0, applied.stdout + applied.stderr)
        self.assertEqual(payload["verify"]["new_comments"], 1)
        _, inspected = self.run_json(REDLINE, "inspect", "--source", str(source), "--limit", "0")
        self.assertEqual(inspected["paragraphs"][0]["text"], "付款期限为30日。")
        self.assertEqual(inspected["paragraphs"][0]["rejected_text"], "付款期限为30日。")

    def test_render_engine_rejects_mismatched_block_anchor_without_fuzzy_search(self) -> None:
        source = self.make_docx("合同.docx", ["甲方应付款。"])
        snapshot = self.snapshot(source)
        block = snapshot["paragraphs"][0]
        plan = {
            "source": str(source),
            "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
            "operations": [
                {
                    "type": "render_block",
                    "locator": {"paragraph_index": 1, "block_id": block["block_id"]},
                    "original_text": "甲方须付款。",
                    "original_text_sha256": hashlib.sha256("甲方须付款。".encode()).hexdigest(),
                    "new_text": "甲方应及时付款。",
                    "intent_id": "PAY-01",
                    "dependency_group": "PAY",
                }
            ],
        }
        original_bytes = source.read_bytes()
        result, error = self.run_json(REDLINE, "apply", payload=plan)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(error["error_code"], "BLOCK_ANCHOR_MISMATCH")
        self.assertEqual(source.read_bytes(), original_bytes)

    def test_unrelated_failed_group_is_isolated_and_evidence_is_preserved(self) -> None:
        source = self.make_mixed_format_docx("合同.docx")
        snapshot = self.snapshot(source)
        intents = [
            self.edit_intent(snapshot, 1, "可安全修改的付款条款，并明确期限。", intent_id="GOOD"),
            self.edit_intent(snapshot, 2, "混合格式条款，并补充责任。", intent_id="BAD"),
        ]
        plan_path = self.evidence_root / "partial.json"
        evidence = self.evidence_root / "partial"
        result, prepared = self.run_json(
            PIPELINE,
            "prepare",
            "--output-plan",
            str(plan_path),
            "--evidence-dir",
            str(evidence),
            payload=self.intent_payload(source, snapshot, intents),
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(prepared["stage"], "prepared_partial")
        self.assertEqual(prepared["accepted_group_count"], 1)
        self.assertEqual(prepared["rejected_group_count"], 1)
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        self.assertEqual([op["intent_id"] for op in plan["operations"]], ["GOOD"])
        group_evidence = json.loads((evidence / "03-group-results.json").read_text(encoding="utf-8"))
        self.assertEqual({item["status"] for item in group_evidence["groups"]}, {"accepted", "rejected"})

    def test_plan_file_is_rejected_inside_contract_directory(self) -> None:
        source = self.make_docx("合同.docx", ["甲方付款。"])
        snapshot = self.snapshot(source)
        intent = self.edit_intent(snapshot, 1, "甲方应及时付款。", intent_id="PAY-01")
        result, error = self.run_json(
            PIPELINE,
            "prepare",
            "--output-plan",
            str(self.root / "forbidden-plan.json"),
            payload=self.intent_payload(source, snapshot, [intent]),
        )
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(error["error_code"], "INVALID_PLAN_PATH")

    def test_dependency_group_is_atomic(self) -> None:
        source = self.make_mixed_format_docx("合同.docx")
        snapshot = self.snapshot(source)
        intents = [
            self.edit_intent(
                snapshot,
                1,
                "可安全修改的付款条款，并明确期限。",
                intent_id="LINKED-1",
                dependency_group="LINKED",
            ),
            self.edit_intent(
                snapshot,
                2,
                "混合格式条款，并补充责任。",
                intent_id="LINKED-2",
                dependency_group="LINKED",
            ),
            self.edit_intent(snapshot, 3, "另一安全条款，并补充通知。", intent_id="INDEPENDENT"),
        ]
        plan_path = self.evidence_root / "grouped.json"
        result, prepared = self.run_json(
            PIPELINE,
            "prepare",
            "--output-plan",
            str(plan_path),
            payload=self.intent_payload(source, snapshot, intents),
        )
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(prepared["stage"], "prepared_partial")
        plan = json.loads(plan_path.read_text(encoding="utf-8"))
        self.assertEqual([op["intent_id"] for op in plan["operations"]], ["INDEPENDENT"])


if __name__ == "__main__":
    unittest.main()
